# -*- coding: utf-8 -*-
#
# -----------------------------------------------------------------------------
# Nome do Arquivo: Buscar.py
#
# Autor:          [Aldair José]
# E-mail:         [aldairsjr@gmail.com]
# Data:           [Data de Criação: 03/06/2026]
# Versão:         [2.3.3Métricas]
# -----------------------------------------------------------------------------

import re
import time
import os
from datetime import datetime
import pandas as pd
import requests
import subprocess
import undetected_chromedriver as uc
from bs4 import BeautifulSoup
from selenium.webdriver.common.by import By
from selenium.common.exceptions import NoSuchElementException
import openpyxl
from openpyxl import load_workbook, Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side

URL_BASE_TRF5 = "https://cp.trf5.jus.br/cp/cp.do"
HEADERS = {"Content-Type": "application/x-www-form-urlencoded"}
METRICAS_FILE = 'Métricas de Tempo.xlsx'

def registrar_metrica(tipo, identificador, tempo_gasto):
    """
    Registra o tempo gasto para busca e extração em um arquivo Excel separado.
    """
    """Definição de estilos profissionais"""
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    summary_hdr_fill = PatternFill(start_color="244062", end_color="244062", fill_type="solid")
    zebra_fill = PatternFill(start_color="F2F5F8", end_color="F2F5F8", fill_type="solid")
    
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=11, bold=True)
    font_normal = Font(name="Calibri", size=11)
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    center_align = Alignment(horizontal='center', vertical='center')
    left_align = Alignment(horizontal='left', vertical='center')
    
    aba_metrica = tipo[:31]
    
    if not os.path.exists(METRICAS_FILE):
        workbook = openpyxl.Workbook()
        default_sheet = workbook.active
        workbook.remove(default_sheet)
    else:
        try:
            workbook = openpyxl.load_workbook(METRICAS_FILE)
        except Exception:
            workbook = openpyxl.Workbook()
            default_sheet = workbook.active
            workbook.remove(default_sheet)

    if aba_metrica not in workbook.sheetnames:
        sheet = workbook.create_sheet(title=aba_metrica)
        sheet.views.sheetView[0].showGridLines = True
        
        """Cabeçalhos do Log Detalhado"""
        headers = ["Data/Hora", "Identificador", "Tempo (segundos)"]
        for col_idx, h in enumerate(headers, 1):
            cell = sheet.cell(row=1, column=col_idx, value=h)
            cell.font = font_header
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border
            
        """Cabeçalhos da Tabela de Resumo Estatístico (Colunas E e F)"""
        sheet.cell(row=1, column=5, value="Métrica Resumo").font = font_header
        sheet.cell(row=1, column=5).fill = summary_hdr_fill
        sheet.cell(row=1, column=5).alignment = left_align
        sheet.cell(row=1, column=5).border = thin_border
        
        sheet.cell(row=1, column=6, value="Valor").font = font_header
        sheet.cell(row=1, column=6).fill = summary_hdr_fill
        sheet.cell(row=1, column=6).alignment = center_align
        sheet.cell(row=1, column=6).border = thin_border
        
        indicadores = ["Total de Quantidade", "Média de Tempo (s)"]
        for idx, ind in enumerate(indicadores, 2):
            c_ind = sheet.cell(row=idx, column=5, value=ind)
            c_ind.font = font_bold
            c_ind.border = thin_border
            c_ind.alignment = left_align
            
            c_val = sheet.cell(row=idx, column=6)
            c_val.font = font_bold
            c_val.border = thin_border
            c_val.alignment = center_align
    else:
        sheet = workbook[aba_metrica]

    """Localizar a próxima linha disponível no log"""
    proxima_linha = 2
    while sheet.cell(row=proxima_linha, column=1).value is not None:
        proxima_linha += 1

    """Gravar as informações do item individual"""
    data_atual = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    valores_log = [data_atual, identificador, round(tempo_gasto, 2)]
    
    for col_idx, val in enumerate(valores_log, 1):
        cell = sheet.cell(row=proxima_linha, column=col_idx, value=val)
        cell.font = font_normal
        cell.alignment = center_align
        cell.border = thin_border
        if proxima_linha % 2 == 0:
            cell.fill = zebra_fill

    """Atualizar as fórmulas do painel de resumo utilizando referências do Excel"""
    sheet.cell(row=2, column=6, value=f"=COUNTA(A2:A{proxima_linha})").number_format = '#,##0'
    sheet.cell(row=3, column=6, value=f"=AVERAGE(C2:C{proxima_linha})").number_format = '0.00'

    """Ajustar as larguras das colunas"""
    larguras_colunas = {'A': 22, 'B': 25, 'C': 18, 'D': 4, 'E': 22, 'F': 15}
    for col, largura in larguras_colunas.items():
        sheet.column_dimensions[col].width = largura

    try:
        workbook.save(METRICAS_FILE)
    except Exception as e:
        print(f"Erro ao salvar arquivo de métricas de tempo: {e}")

def formatar_numero_limpo(numero_original):
    if not isinstance(numero_original, str):
        numero_original = str(numero_original)
    return numero_original.replace(".", "").replace("-", "").replace("/", "").replace(" ", "")

def formatar_requisitorio_17_digitos(numero_original):
    numero_limpo = formatar_numero_limpo(numero_original)
    if len(numero_limpo) == 17:
        n = numero_limpo
        return f"{n[0:4]}.{n[4:6]}.{n[6:8]}.{n[8:11]}.{n[11:17]}"
    return numero_original

def formatar_npu_20_digitos(numero_original):
    numero_limpo = formatar_numero_limpo(numero_original)
    if len(numero_limpo) == 20:
        n = numero_limpo
        return f"{n[0:7]}-{n[7:9]}.{n[9:13]}.{n[13:14]}.{n[14:16]}.{n[16:20]}"
    return numero_original

class Precatorio:
    def __init__(self, requisitorio):
        self.requisitorio_original = requisitorio
        self.html_content = self._obter_html_extrato(requisitorio)

        self.numero_requisitorio = ""
        self.processo_execucao = ""
        self.vara = ""
        self.requerente = ""
        self.requerido = ""
        self.banco = ""
        self.data_individualizacao = None
        self.data_pgto = None
        self.data_cancelamento = None
        self.precatorio_numero = ""
        self.processo_principal_npu = ""
        self.advogados = []
        self.escritorios = []
        self.deprecantes = []

        if self.html_content:
            self._extrair_dados()

    def is_localizado(self):
        return bool(self.numero_requisitorio)

    def _obter_html_extrato(self, requisitorio):
        payload = {"tipo": "xmlproc", "ordenacao+cpf": "D", "tiproc": "T", "filtro": requisitorio,
                   "Pesquisar": "Pesquisar"}
        try:
            response = requests.post(URL_BASE_TRF5, data=payload, headers=HEADERS, timeout=10)
            response.raise_for_status()
            response.encoding = 'utf-8'
            if "\nNenhum resultado encontrado" in response.text:
                return None
            return response.text
        except requests.exceptions.RequestException as e:
            print(f"\nErro ao buscar requisitório {requisitorio}")
            return None

    def _extrair_com_regex(self, pattern):
        if not self.html_content: return ""
        match = re.search(pattern, self.html_content, re.IGNORECASE | re.DOTALL)
        if match:
            texto_cru = match.group(1)
            return BeautifulSoup(texto_cru, "html.parser").get_text(strip=True)
        return ""

    def _extrair_multiplos_com_regex(self, pattern):
        """Função adicionada para capturar todas as ocorrências de um padrão"""
        if not self.html_content: return []
        matches = re.finditer(pattern, self.html_content, re.IGNORECASE | re.DOTALL)
        resultados = []
        for match in matches:
            texto_cru = match.group(1)
            resultados.append(BeautifulSoup(texto_cru, "html.parser").get_text(strip=True))
        return resultados

    def _extrair_dados(self):
        self.processo_principal_npu = self._extrair_com_regex(r'PROCESSO N(?:&ordm;|º)\s*([\d.-]+)')
        self.precatorio_numero = self._extrair_com_regex(
            r'(?:PRECAT(?:&Oacute;|Ó)RIO|REQUISI(?:&Ccedil;|Ç)AO DE PEQUENO VALOR)\s*\((.*?)-')
        self.numero_requisitorio = self._extrair_com_regex(r'N(?:&Uacute;|Ú)MERO DO REQUISIT(?:&Oacute;|Ó)RIO:\s*([^<]+)')
        self.processo_execucao = self._extrair_com_regex(
            r'N(?:&Uacute;|Ú)MERO DO PROCESSO DE EXECU(?:&Ccedil;|&Ccedil;|Ç)(?:&Atilde;|Ã)O:\s*([^<]+)')
        self.vara = self._extrair_com_regex(r'VARA:\s*([^<]+)')
        self.requerente = self._extrair_com_regex(r'REQTE<\/td>.*?<b>(.*?)<\/b>')
        self.requerido = self._extrair_com_regex(r'REQDO<\/td>.*?<b>(.*?)<\/b>')
        
        """Chamadas alteradas para capturar múltiplos registros"""
        self.advogados = self._extrair_multiplos_com_regex(r'Advogado/Procurador:<\/b><\/td>\s*<td.*?>(.*?)<\/td>')
        self.escritorios = self._extrair_multiplos_com_regex(
            r'Escrit(?:&oacute;|ó)rio Advocat(?:&iacute;|í)cio:<\/b><\/td>\s*<td.*?>(.*?)<\/td>')
        self.deprecantes = self._extrair_multiplos_com_regex(r'DEPRECANTE:<\/b><\/td>\s*<td.*?>(.*?)<\/td>')

        padrao_banco = r'Institui(?:&ccedil;?|ç)(?:&atilde;?|ã)o Financeira para Pagamento Banco:\s*([^<]+)'
        match_banco = re.search(padrao_banco, self.html_content, re.IGNORECASE | re.DOTALL)
        if match_banco:
            texto_completo_banco = BeautifulSoup(match_banco.group(1), "html.parser").get_text(strip=True)
            self.banco = texto_completo_banco.split('-')[0].strip() if '-' in texto_completo_banco else texto_completo_banco

        data_ind_str = self._extrair_com_regex(r'name="mov_0">Em (\d{2}\/\d{2}\/\d{4})')
        self.data_individualizacao = pd.to_datetime(data_ind_str, format='%d/%m/%Y', errors='coerce')

        soup = BeautifulSoup(self.html_content, 'html.parser')
        self.data_pgto = pd.NaT
        self.data_cancelamento = pd.NaT

        movimento_anchors = soup.find_all('a', attrs={'name': re.compile(r'^mov_\d+$')})

        for anchor in movimento_anchors:
            tabela_movimento = anchor.find_parent('table')
            if not tabela_movimento:
                continue

            texto_tabela = tabela_movimento.get_text().lower()

            if pd.isna(self.data_pgto) and "depósito em conta" in texto_tabela:
                texto_data = anchor.get_text()
                match_data = re.search(r'(\d{2}/\d{2}/\d{4})', texto_data)
                if match_data:
                    self.data_pgto = pd.to_datetime(match_data.group(1), format='%d/%m/%Y', errors='coerce')

            if pd.isna(self.data_cancelamento) and "cancelamento" in texto_tabela:
                texto_data = anchor.get_text()
                match_data = re.search(r'(\d{2}/\d{2}/\d{4})', texto_data)
                if match_data:
                    self.data_cancelamento = pd.to_datetime(match_data.group(1), format='%d/%m/%Y', errors='coerce')

            if not pd.isna(self.data_pgto) and not pd.isna(self.data_cancelamento):
                break

    def to_dict(self):
        dados = {
            "PRC Pesquisado": self.requisitorio_original, 
            "PROCESSO Nº": self.processo_principal_npu,
            "PRECATÓRIO/RPV": self.precatorio_numero,
            "NÚMERO DO REQUISITÓRIO": formatar_requisitorio_17_digitos(self.numero_requisitorio),
            "NÚMERO DO PROCESSO DE EXECUÇÃO": formatar_npu_20_digitos(self.processo_execucao), 
            "VARA": self.vara,
            "REQTE": self.requerente, 
            "REQDO": self.requerido, 
            "Banco": self.banco,
            "Data Individualização": self.data_individualizacao, 
            "Data Pagamento": self.data_pgto,
            "Data Cancelamento": self.data_cancelamento,
        }

        """Tratamento de múltiplos Advogados/Procuradores"""
        if self.advogados:
            dados["Advogado/Procurador"] = self.advogados[0]
            for i in range(1, len(self.advogados)):
                dados[f"{i+1}º Advogado/Procurador"] = self.advogados[i]
        else:
            dados["Advogado/Procurador"] = ""

        """Tratamento de múltiplos Escritórios Advocatícios"""
        if self.escritorios:
            dados["Escritório Advocatício"] = self.escritorios[0]
            for i in range(1, len(self.escritorios)):
                dados[f"{i+1}º Escritório Advocatício"] = self.escritorios[i]
        else:
            dados["Escritório Advocatício"] = ""

        """Tratamento de múltiplos Deprecantes (geralmente há 1, mas cobrimos mais de 1 caso exista)"""
        if self.deprecantes:
            dados["Deprecante"] = self.deprecantes[0]
            for i in range(1, len(self.deprecantes)):
                dados[f"{i+1}º Deprecante"] = self.deprecantes[i]
        else:
            dados["Deprecante"] = ""

        return dados

def pesquisar_por_prc(lista_prc):
    print(f"\nIniciando busca por PRC...\n")
    resultados, status_list = [], []
    ordem_colunas = ["PRC Pesquisado", "PROCESSO Nº", "PRECATÓRIO/RPV", "NÚMERO DO REQUISITÓRIO",
                     "NÚMERO DO PROCESSO DE EXECUÇÃO", "VARA", "REQTE", "REQDO", "Banco", "Data Individualização",
                     "Data Pagamento", "Data Cancelamento", "Advogado/Procurador", "Escritório Advocatício",
                     "Deprecante"]

    for i, prc_num in enumerate(lista_prc):
        inicio_busca = time.time()
        
        prc_base = prc_num.split('-')[0]
        numero_formatado = formatar_numero_limpo(prc_base)

        print(f"Processando {i + 1}/{len(lista_prc)}: {prc_num} (buscando por: {numero_formatado})")

        precatorio = Precatorio(numero_formatado)
        if precatorio.is_localizado():
            status_list.append("Localizado")
            dados = precatorio.to_dict()
            dados["PRC Pesquisado"] = prc_num
            resultados.append(dados)
        else:
            status_list.append("Não localizado")
        
        fim_busca = time.time()
        registrar_metrica("Pesquisa PRC", prc_num, fim_busca - inicio_busca)
            
        time.sleep(1)

    df = pd.DataFrame(resultados)
    if df.empty:
        return pd.DataFrame(columns=ordem_colunas), status_list

    """Garantir que colunas adicionadas dinamicamente sejam incluídas na ordem correta"""
    todas_colunas = df.columns.tolist()
    colunas_finais = []
    
    for col in ordem_colunas:
        if col in todas_colunas:
            colunas_finais.append(col)
            """Verifica se existem colunas dinâmicas relacionadas (ex: "2º Advogado/Procurador")"""
            if col in ["Advogado/Procurador", "Escritório Advocatício", "Deprecante"]:
                extra_cols = []
                for c in todas_colunas:
                    if c.endswith(f"º {col}"):
                        try:
                            num = int(c.split("º")[0])
                            extra_cols.append((num, c))
                        except ValueError:
                            pass
                extra_cols.sort()
                colunas_finais.extend([c[1] for c in extra_cols])

    return df[colunas_finais], status_list


def pesquisar_prc_por_requisitorio(requisitorio: str) -> str:
    url = "https://cp.trf5.jus.br/cp/cp.do"
    payload = {
        "tipo": "xmlrpvprec", "tipoproc": "T",
        "numRequisitorio": requisitorio, "ordenacao": "D"
    }
    headers = {
        "Content-Type": "application/x-www-form-urlencoded",
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/91.0.4472.124 Safari/537.36"
    }
    try:
        pagina = requests.post(url, data=payload, headers=headers, timeout=10)
        pagina.raise_for_status()
        soup = BeautifulSoup(pagina.text, 'lxml')
        tabela_resultados = soup.find('table', class_='consulta_resultados')
        if not tabela_resultados:
            return "Não encontrado"

        linhas = tabela_resultados.find_all('tr')
        if len(linhas) > 1:
            celulas = linhas[1].find_all('td')
            if len(celulas) > 2:
                prc_bruto = celulas[2].get_text(strip=True)
                prc_limpo = prc_bruto.replace("-AL (@)", "").strip()
                return prc_limpo
        return "Tabela não reconhecida"
    except requests.exceptions.RequestException:
        return "Erro de conexão"


def atualizar_requisitorios_na_planilha(caminho_arquivo: str, nome_aba: str):
    print(f"\nIniciando busca por Requisitório e atualização da planilha...\n")

    coluna_leitura = "Requisitório"
    coluna_escrita = "PRC/RPV"
    center_align = Alignment(horizontal='center', vertical='center')

    try:
        workbook = load_workbook(caminho_arquivo)
        if nome_aba not in workbook.sheetnames:
            print(f"Erro: Aba '{nome_aba}' não encontrada no arquivo '{caminho_arquivo}'.")
            return
        sheet = workbook[nome_aba]
    except FileNotFoundError:
        print(f"Erro: O arquivo '{caminho_arquivo}' não foi encontrado.")
        return
    except Exception as e:
        print(f"Erro ao carregar o arquivo: {e}")
        return

    headers = [cell.value for cell in sheet[1]]
    try:
        col_leitura_idx = headers.index(coluna_leitura) + 1
    except ValueError:
        print(f"Erro: A coluna de leitura '{coluna_leitura}' não foi encontrada na aba '{nome_aba}'.")
        return

    if coluna_escrita not in headers:
        col_escrita_idx = len(headers) + 1
        header_cell = sheet.cell(row=1, column=col_escrita_idx, value=coluna_escrita)
        header_cell.alignment = center_align
        print(f"Coluna '{coluna_escrita}' não encontrada. Foi criada na posição {col_escrita_idx}.")
    else:
        col_escrita_idx = headers.index(coluna_escrita) + 1

    sheet.column_dimensions[get_column_letter(col_escrita_idx)].width = 20

    total_rows = sheet.max_row
    for i, row in enumerate(sheet.iter_rows(min_row=2, max_row=total_rows), start=2):
        requisitorio_cell = row[col_leitura_idx - 1]
        resultado_cell = row[col_escrita_idx - 1]

        requisitorio_val = requisitorio_cell.value

        if not requisitorio_val or resultado_cell.value is not None:
            continue
            
        inicio_busca = time.time()

        requisitorio_str = formatar_numero_limpo(str(requisitorio_val))
        print(f"Processando {i - 1}/{total_rows - 1}: {requisitorio_str}")

        resultado = pesquisar_prc_por_requisitorio(requisitorio_str)
        print(f"PRC/RPV encontrado: {resultado}")

        resultado_cell.value = resultado
        resultado_cell.alignment = center_align

        try:
            workbook.save(caminho_arquivo)
            print(f"Progresso salvo\n")
        except Exception as e:
            print(
                f"\nERRO AO SALVAR: Não foi possível salvar a alteração para a linha {i}. Tente fechar o arquivo Excel se estiver aberto.")
            time.sleep(5)
            
        fim_busca = time.time()
        registrar_metrica("Pesquisa Requisitório", requisitorio_str, fim_busca - inicio_busca)

        time.sleep(1)


def _obter_versao_chrome_windows(caminho_executavel):
    try:
        if not os.path.exists(caminho_executavel):
            return None
        
        cmd = f'wmic datafile where name="{caminho_executavel.replace("\\", "\\\\")}" get Version /value'
        resultado = subprocess.check_output(cmd, shell=True).decode()

        versao_completa = resultado.strip().split('=')[1]
        versao_principal = int(versao_completa.split('.')[0])
        return versao_principal
    except Exception:
        return None


def _pesquisar_com_selenium(lista_de_itens, tipo_pesquisa, id_radio_button, id_input_field):
    print(f"\nIniciando busca por {tipo_pesquisa}...")

    os.environ['WDM_LOG_LEVEL'] = '0'
    options = uc.ChromeOptions()
    options.add_argument("--headless")
    options.add_argument("--log-level=3")
    options.add_argument("--disable-gpu")
    options.add_argument("--no-sandbox")

    """Tentativa de localizar o Chrome no Windows"""
    caminho_chrome = None
    versao_main = None
    caminhos_possiveis = [
        r"C:\Program Files\Google\Chrome\Application\chrome.exe",
        r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
        os.path.expanduser(r"~\AppData\Local\Google\Chrome\Application\chrome.exe")
    ]

    for caminho in caminhos_possiveis:
        if os.path.exists(caminho):
            caminho_chrome = caminho
            break

    if caminho_chrome:
        versao_main = _obter_versao_chrome_windows(caminho_chrome)

    try:
        if versao_main and caminho_chrome:
            print(f"Versão do Chrome detectado: {versao_main}")
            
            driver = uc.Chrome(
                options=options,
                version_main=versao_main,
                browser_executable_path=caminho_chrome
            )
        else:
            print("Chrome não detectado automaticamente nos caminhos padrão. Tentando inicialização padrão.")
            driver = uc.Chrome(options=options)
            
    except Exception as e:
        print("\n" + "=" * 50)
        print(f"ERRO CRÍTICO: Não foi possível iniciar o Undetected Chrome Driver.")
        print(f"Detalhe do erro: {e}")
        print("=" * 50 + "\n")
        return pd.DataFrame(), ["Erro de Driver"] * len(lista_de_itens)

    coluna_pesquisada = f"{tipo_pesquisa} Pesquisado"
    resultados_finais, status_list = [], []
    ordem_colunas = [
        coluna_pesquisada, "Nº Processo no TRF", "PRC/RPV", "REQTE", "Advogado/Procurador",
        "Escritório Advocatício", "REQDO", "Deprecante", "Data do Movimento", "Hora do Movimento", "Última Movimentação"
    ]

    for i, item_num in enumerate(lista_de_itens):
        inicio_busca = time.time()
        item_formatado = formatar_numero_limpo(item_num)
        print(f"Processando {i + 1}/{len(lista_de_itens)}: {item_num}")

        try:
            driver.get("https://cp.trf5.jus.br/cp/index.html")
            driver.find_element(By.ID, id_radio_button).click()
            driver.find_element(By.ID, id_input_field).clear()
            driver.find_element(By.ID, id_input_field).send_keys(item_formatado)
            driver.find_element(By.ID, "submitConsulta").click()
            time.sleep(2.5)

            if len(driver.window_handles) < 2:
                status_list.append("Não localizado")
                resultados_finais.append({coluna_pesquisada: item_num, "Processo": "Nenhum resultado encontrado"})
                fim_busca = time.time()
                registrar_metrica(f"Pesquisa {tipo_pesquisa}", item_num, fim_busca - inicio_busca)
                continue

            driver.switch_to.window(driver.window_handles[-1])

            foi_localizado_para_este_item = False
            dados_processo_atual = {}

            while True:
                soup = BeautifulSoup(driver.page_source, 'html.parser')
                tabela = soup.find('table', class_='consulta_resultados')

                if tabela and len(tabela.find_all('tr')) > 1:
                    if not foi_localizado_para_este_item:
                        foi_localizado_para_este_item = True

                    header_map = {cell.get_text(strip=True): idx for idx, cell in
                                  enumerate(tabela.find('tr').find_all(['th', 'td'])) if cell.get_text(strip=True)}

                    for linha in tabela.find_all('tr')[1:]:
                        celulas_texto = [td.get_text(strip=True) for td in linha.find_all('td')]
                        if not celulas_texto: continue
                        idx_processo = header_map.get("Nº Processo no TRF") or header_map.get("Processo")

                        if idx_processo is not None and len(celulas_texto) > idx_processo and re.match(r'[\d.-]{15,}', celulas_texto[idx_processo]):
                            if dados_processo_atual:
                                resultados_finais.append(dados_processo_atual)
                            dados_processo_atual = {coluna_pesquisada: item_num}
                            for header_name, idx in header_map.items():
                                if len(celulas_texto) > idx:
                                    dados_processo_atual[header_name] = celulas_texto[idx]
                            if len(celulas_texto) > 3 and 'REQTE' not in dados_processo_atual and celulas_texto[3] not in dados_processo_atual.values():
                                dados_processo_atual['REQTE'] = celulas_texto[3]
                        elif len(celulas_texto) > 2:
                            for j in range(1, len(celulas_texto), 2):
                                if (j + 1) < len(celulas_texto):
                                    chave = celulas_texto[j].replace(':', '').strip()
                                    valor = celulas_texto[j + 1]
                                    if chave in ordem_colunas and dados_processo_atual:
                                        dados_processo_atual[chave] = valor

                try:
                    next_page_button = driver.find_element(By.LINK_TEXT, ">")
                    next_page_button.click()
                    time.sleep(2.5)
                except NoSuchElementException:
                    break

            if dados_processo_atual:
                resultados_finais.append(dados_processo_atual)

            if foi_localizado_para_este_item:
                status_list.append("Localizado")
            else:
                status_list.append("Não localizado")
                resultados_finais.append({coluna_pesquisada: item_num, "Processo": "Nenhum resultado encontrado"})

            if len(driver.window_handles) > 1:
                driver.close()
                driver.switch_to.window(driver.window_handles[0])

        except Exception as e:
            print(f"Ocorreu um erro ao processar o item {item_num}: {e}")
            status_list.append("Erro na busca")
            if len(driver.window_handles) > 1:
                driver.close()
                driver.switch_to.window(driver.window_handles[0])
                
        fim_busca = time.time()
        registrar_metrica(f"Pesquisa {tipo_pesquisa}", item_num, fim_busca - inicio_busca)

    driver.quit()

    df = pd.DataFrame(resultados_finais)
    if df.empty:
        return pd.DataFrame(columns=ordem_colunas), status_list

    df = df.rename(columns={"Processo": "Nº Processo no TRF", "Classe": "PRC/RPV"})

    for col in ordem_colunas:
        if col not in df.columns:
            df[col] = ""
    return df[ordem_colunas], status_list


def pesquisar_por_cpf(lista_cpf):
    return _pesquisar_com_selenium(lista_cpf, "CPF", "tipo_xmlcpf", "filtroCPF2")


def pesquisar_por_originario(lista_originarios):
    return _pesquisar_com_selenium(lista_originarios, "Originário", "tipo_xmlrpvprec", "numOriginario")


def pesquisar_por_execucao(lista_execucoes):
    return _pesquisar_com_selenium(lista_execucoes, "Execução", "tipo_xmlrpvprec", "numProcessExec")


def pesquisar_por_npu(lista_npu):
    print(f"\nIniciando busca por NPU...\n")
    resultados, status_list = [], []
    ordem_colunas = ["NPU Pesquisado", "PROCESSO Nº", "PRECATÓRIO/RPV", "NÚMERO DO REQUISITÓRIO",
                     "NÚMERO DO PROCESSO DE EXECUÇÃO", "VARA", "REQTE", "REQDO", "Banco", "Data Individualização",
                     "Data Pagamento", "Data Cancelamento", "Advogado/Procurador", "Escritório Advocatício",
                     "Deprecante"]
    for i, npu_num in enumerate(lista_npu):
        inicio_busca = time.time()
        
        numero_formatado = formatar_numero_limpo(npu_num)
        print(f"Processando {i + 1}/{len(lista_npu)}: {npu_num}")
        precatorio = Precatorio(numero_formatado)
        if precatorio.is_localizado():
            status_list.append("Localizado")
            dados = precatorio.to_dict()
            dados["NPU Pesquisado"] = npu_num
            dados.pop("PRC Pesquisado", None)
            resultados.append(dados)
        else:
            status_list.append("Não localizado")
            
        fim_busca = time.time()
        registrar_metrica("Pesquisa NPU", npu_num, fim_busca - inicio_busca)
        time.sleep(1)
        
    df = pd.DataFrame(resultados)
    if df.empty:
        return pd.DataFrame(columns=ordem_colunas), status_list
    
    todas_colunas = df.columns.tolist()
    colunas_finais = []
    
    for col in ordem_colunas:
        if col in todas_colunas:
            colunas_finais.append(col)
            if col in ["Advogado/Procurador", "Escritório Advocatício", "Deprecante"]:
                extra_cols = []
                for c in todas_colunas:
                    if c.endswith(f"º {col}"):
                        try:
                            num = int(c.split("º")[0])
                            extra_cols.append((num, c))
                        except ValueError:
                            pass
                extra_cols.sort()
                colunas_finais.extend([c[1] for c in extra_cols])

    return df[colunas_finais], status_list


def atualizar_aba_com_formatacao(df, nome_arquivo, nome_aba):
    center_align = Alignment(horizontal='center', vertical='center')
    try:
        book = load_workbook(nome_arquivo)
    except FileNotFoundError:
        book = Workbook()
        if 'Sheet' in book.sheetnames:
            book.remove(book['Sheet'])
        print(f"Arquivo '{nome_arquivo}' não encontrado. Um novo será criado")

    if nome_aba in book.sheetnames:
        sheet = book[nome_aba]
        for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
            for cell in row:
                cell.value = None

        start_row = 2
        for r_idx, row_data in enumerate(dataframe_to_rows(df, index=False, header=False), start=start_row):
            for c_idx, value in enumerate(row_data, start=1):
                cell = sheet.cell(row=r_idx, column=c_idx, value=value)
                cell.alignment = center_align
    else:
        sheet = book.create_sheet(nome_aba)
        for r_idx, row_data in enumerate(dataframe_to_rows(df, index=False, header=True), start=1):
            for c_idx, value in enumerate(row_data, start=1):
                cell = sheet.cell(row=r_idx, column=c_idx, value=value)
                cell.alignment = center_align

    try:
        book.save(nome_arquivo)
    except Exception as e:
        print(f"Ocorreu um erro ao salvar o arquivo '{nome_arquivo}': {e}")


if __name__ == "__main__":
    arquivo_entrada = 'Entrada.xlsx'
    arquivo_saida = 'Resultados.xlsx'
    while True:

        print("1 - Pesquisar por PRC/RPV")
        print("2 - Pesquisar por NPU")
        print("3 - Pesquisar por CPF")
        print("4 - Pesquisar por Processo Originário")
        print("5 - Pesquisar por Processo de Execução")
        print("6 - Pesquisar por Requisitório")
        print("")
        print("7 - Sair")

        escolha = input("\nDigite a opção desejada: ")

        if escolha == '1':
            try:
                df_entrada_prc = pd.read_excel(arquivo_entrada, sheet_name='PRCs', dtype={'PRC/RPV': str})
                lista_de_prcs = df_entrada_prc['PRC/RPV'].dropna().tolist()
                if lista_de_prcs:
                    df_resultados_prc, status_prc = pesquisar_por_prc(lista_de_prcs)

                    df_entrada_prc['Status'] = status_prc

                    colunas_data = ["Data Individualização", "Data Pagamento", "Data Cancelamento"]
                    for col in colunas_data:
                        if col in df_resultados_prc.columns:
                            df_resultados_prc[col] = pd.to_datetime(df_resultados_prc[col], errors='coerce').dt.strftime('%d/%m/%Y').replace('NaT', '')

                    atualizar_aba_com_formatacao(df_entrada_prc, arquivo_entrada, 'PRCs')
                    atualizar_aba_com_formatacao(df_resultados_prc, arquivo_saida, 'Resultados PRC ou RPV')
                    print(f"\nDados salvos na aba Resultados PRC ou RPV\n")
                else:
                    print("\nNenhum PRC encontrado na coluna PRC/RPV\n")
            except FileNotFoundError:
                print(f"Erro: Aba PRCs não encontrada no arquivo {arquivo_entrada}")
            except Exception as e:
                print(f"Ocorreu um erro: {e}")

        elif escolha == '2':
            try:
                df_entrada_npu = pd.read_excel(arquivo_entrada, sheet_name='NPUs', dtype={'NPU': str})
                lista_de_npus = df_entrada_npu['NPU'].dropna().tolist()
                if lista_de_npus:
                    df_resultados_npu, status_npu = pesquisar_por_npu(lista_de_npus)
                    df_entrada_npu['Status'] = status_npu
                    colunas_data = ["Data Individualização", "Data Pagamento", "Data Cancelamento"]
                    for col in colunas_data:
                        if col in df_resultados_npu.columns:
                            df_resultados_npu[col] = pd.to_datetime(df_resultados_npu[col], errors='coerce').dt.strftime('%d/%m/%Y').replace('NaT', '')

                    atualizar_aba_com_formatacao(df_entrada_npu, arquivo_entrada, 'NPUs')
                    atualizar_aba_com_formatacao(df_resultados_npu, arquivo_saida, 'Resultados NPU')
                    print(f"\nDados salvos na aba Resultados NPU\n")
                else:
                    print("\nNenhum NPU encontrado na coluna NPU\n")
            except FileNotFoundError:
                print(f"Erro: Aba NPUs não encontrada no arquivo {arquivo_entrada}")
            except Exception as e:
                print(f"Ocorreu um erro: {e}")

        elif escolha == '3':
            try:
                df_entrada_cpf = pd.read_excel(arquivo_entrada, sheet_name='CPFs', dtype={'CPF': str})
                lista_de_cpfs = df_entrada_cpf['CPF'].dropna().tolist()
                if lista_de_cpfs:
                    df_resultados_cpf, status_cpf = pesquisar_por_cpf(lista_de_cpfs)
                    df_entrada_cpf['Status'] = status_cpf
                    atualizar_aba_com_formatacao(df_entrada_cpf, arquivo_entrada, 'CPFs')
                    atualizar_aba_com_formatacao(df_resultados_cpf, arquivo_saida, 'Resultados CPF')
                    print(f"\nDados salvos na aba Resultados CPF\n")
                else:
                    print(f"\nNenhum CPF encontrado na coluna CPF\n")
            except FileNotFoundError:
                print(f"Erro: Aba CPFs não encontrada no arquivo {arquivo_entrada}")
            except Exception as e:
                print(f"Ocorreu um erro: {e}")

        elif escolha == '4':
            try:
                df_entrada_org = pd.read_excel(arquivo_entrada, sheet_name='Originários', dtype={'Originario': str})
                lista_de_originarios = df_entrada_org['Originário'].dropna().tolist()
                if lista_de_originarios:
                    df_resultados_org, status_org = pesquisar_por_originario(lista_de_originarios)
                    df_entrada_org['Status'] = status_org
                    atualizar_aba_com_formatacao(df_entrada_org, arquivo_entrada, 'Originários')
                    atualizar_aba_com_formatacao(df_resultados_org, arquivo_saida, 'Resultados Originário')
                    print(f"\nDados salvos na aba Resultados por Originário\n")
                else:
                    print("\nNenhum número encontrado na coluna Originário\n")
            except FileNotFoundError:
                print(f"Erro: Aba Originários não encontrada no arquivo {arquivo_entrada}")
            except Exception as e:
                print(f"Ocorreu um erro: {e}")

        elif escolha == '5':
            try:
                df_entrada_exe = pd.read_excel(arquivo_entrada, sheet_name='Execuções', dtype={'Execucao': str})
                lista_de_execucoes = df_entrada_exe['Execução'].dropna().tolist()
                if lista_de_execucoes:
                    df_resultados_exe, status_exe = pesquisar_por_execucao(lista_de_execucoes)
                    df_entrada_exe['Status'] = status_exe
                    atualizar_aba_com_formatacao(df_entrada_exe, arquivo_entrada, 'Execuções')
                    atualizar_aba_com_formatacao(df_resultados_exe, arquivo_saida, 'Resultados Execução')
                    print(f"\nDados salvos na aba Resultados por Execução\n")
                else:
                    print("\nNenhum número encontrado na coluna Execução\n")
            except FileNotFoundError:
                print(f"Erro: Aba Execuções não encontrada no arquivo {arquivo_entrada}")
            except Exception as e:
                print(f"Ocorreu um erro: {e}")

        elif escolha == '6':
            try:
                atualizar_requisitorios_na_planilha(arquivo_entrada, 'Requisitórios')
            except Exception as e:
                print(f"Ocorreu um erro inesperado durante a atualização da planilha: {e}")

        elif escolha == '7':
            break

        else:
            print("\nOpção inválida. Escolha uma das opções acima.")